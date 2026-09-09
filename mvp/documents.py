"""1단계: PDF를 읽고 실제 페이지 번호와 추출된 글을 함께 보관합니다."""

import hashlib
import threading
import time
from collections import OrderedDict
from dataclasses import dataclass
from io import BytesIO

from pypdf import PdfReader
from pypdf.errors import PyPdfError

MAX_FILE_MB = 20
MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024
MAX_PAGES = 1000
_PARSED = OrderedDict()
_PARSE_LOCK = threading.RLock()


class PdfInputError(ValueError):
    """화면에 그대로 표시해도 되는, 원문을 포함하지 않는 오류입니다."""


@dataclass(frozen=True)
class PdfPage:
    # PDF의 첫 장이 1페이지입니다. 인쇄된 쪽수와 다를 수 있습니다.
    number: int | None
    text: str
    error: str | None = None
    location: str = ""
    section: str = ""
    header: str = ""


@dataclass(frozen=True)
class PdfDocument:
    document_name: str
    pages: tuple[PdfPage, ...]
    source_type: str = "pdf"
    warnings: tuple[str, ...] = ()

    @property
    def text_page_count(self) -> int:
        return sum(bool(page.text) for page in self.pages)

    @property
    def character_count(self) -> int:
        return sum(len(page.text) for page in self.pages)


def read_pdf(file_name: str, content: bytes) -> PdfDocument:
    """외부 API 호출이나 파일 저장 없이, 전달받은 PDF 바이트만 읽습니다."""
    if not file_name.lower().endswith(".pdf"):
        raise PdfInputError("PDF 파일을 선택해 주세요. (PDF_TYPE)")
    if not content:
        raise PdfInputError("파일이 비어 있습니다. 원본 파일을 확인해 주세요. (PDF_EMPTY)")
    if len(content) > MAX_FILE_BYTES:
        raise PdfInputError(f"파일은 {MAX_FILE_MB}MB 이하로 선택해 주세요. (PDF_SIZE)")
    if not content.startswith(b"%PDF-"):
        raise PdfInputError("파일 내용이 PDF 형식이 아닙니다. PDF로 다시 저장해 주세요. (PDF_FORMAT)")

    try:
        # 잘못된 파일을 추측해서 복구하지 않고, 읽을 수 없으면 명확히 안내합니다.
        reader = PdfReader(BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise PdfInputError(
                "암호화된 PDF는 이번 단계에서 지원하지 않습니다. "
                "암호가 해제된 사본을 선택해 주세요. (PDF_ENCRYPTED)"
            )
        page_count = len(reader.pages)
        if not page_count:
            raise PdfInputError("페이지가 없는 PDF입니다. 원본 파일을 확인해 주세요. (PDF_NO_PAGES)")
        if page_count > MAX_PAGES:
            raise PdfInputError(
                f"이번 단계에서는 {MAX_PAGES:,}페이지까지 읽습니다. 문서를 나누어 올려 주세요. (PDF_PAGES)"
            )
    except PdfInputError:
        raise
    except (PyPdfError, ValueError, TypeError, KeyError, IndexError, OSError, RecursionError):
        # 파서의 상세 예외에는 원문 일부가 들어갈 수 있어 사용자용 메시지만 제공합니다.
        raise PdfInputError(
            "PDF 구조를 읽을 수 없습니다. 원본을 열어 확인한 뒤 PDF로 다시 저장해 주세요. (PDF_READ)"
        ) from None

    pages = []
    for index in range(page_count):
        try:
            text = reader.pages[index].extract_text() or ""
            # 문장·수치·단위와 줄 순서를 보존합니다. 요약이나 내용 수정은 하지 않습니다.
            text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "").strip()
            pages.append(PdfPage(number=index + 1, text=text))
        except (PyPdfError, ValueError, TypeError, KeyError, IndexError, OSError, RecursionError):
            # 실패한 페이지도 남겨, 이후 페이지의 출처 번호가 당겨지지 않도록 합니다.
            pages.append(
                PdfPage(number=index + 1, text="", error="이 페이지의 텍스트를 추출하지 못했습니다.")
            )

    # 파일명은 표시용으로만 쓰며 디스크 경로로 사용하지 않습니다.
    display_name = file_name.replace("\\", "/").rsplit("/", 1)[-1]
    return PdfDocument(document_name=display_name, pages=tuple(pages))


def read_document(file_name, content, *, ocr=False):
    """서버 메모리에 검토 가능한 추출 결과만 최대 2개/16MB/3분 캐시합니다."""
    from mvp.library import protect_private
    from mvp.pdf_layout import EXTRACTION_VERSION
    key = (file_name, hashlib.sha256(content).hexdigest(), ocr, EXTRACTION_VERSION)
    with _PARSE_LOCK:
        now = time.monotonic()
        for old in list(_PARSED):
            if now - _PARSED[old][0] > 180:
                _PARSED.pop(old)
        if key in _PARSED:
            _PARSED.move_to_end(key)
            return _PARSED[key][1]
    document = _read_document(file_name, content, ocr=ocr)
    protect_private(document.document_name + '\n' + '\n'.join(p.text for p in document.pages))
    size = document.character_count * 4
    if size <= 16_000_000:
        with _PARSE_LOCK:
            _PARSED[key] = (time.monotonic(), document, size)
            while len(_PARSED) > 2 or sum(item[2] for item in _PARSED.values()) > 16_000_000:
                _PARSED.popitem(last=False)
    return document


def _read_document(file_name, content, *, ocr=False):
    """PDF·DOCX·XLSX를 동일한 검색 단위로 변환합니다. Office 쪽수는 추측하지 않습니다."""
    from pathlib import PurePosixPath
    from zipfile import ZipFile

    name = file_name.replace("\\", "/").rsplit("/", 1)[-1]
    kind = PurePosixPath(name).suffix.lower().lstrip(".")
    if kind == "pdf":
        document = read_pdf(name, content)
        from mvp.pdf_layout import enhance_pdf
        document = enhance_pdf(document, content, ocr=ocr)
        if document.character_count > 3_000_000:
            raise PdfInputError("추출된 글이 너무 많습니다. 문서를 나누어 등록하세요. (TEXT_SIZE)")
        return document
    if kind not in {"docx", "xlsx"}:
        raise PdfInputError("PDF, Word(.docx), Excel(.xlsx)를 지원합니다. .doc/.xls는 새 형식으로 저장하세요. (DOCUMENT_TYPE)")
    if not content or len(content) > MAX_FILE_BYTES:
        raise PdfInputError(f"비어 있지 않은 {MAX_FILE_MB}MB 이하 파일을 선택하세요. (DOCUMENT_SIZE)")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if sum(e.file_size for e in entries) > 100 * 1024 * 1024 or len(entries) > 10000:
                raise PdfInputError("압축을 푼 문서가 너무 큽니다. 문서를 나누세요. (OFFICE_SIZE)")
            if any("vbaproject" in e.filename.lower() for e in entries):
                raise PdfInputError("매크로를 제거한 문서를 등록하세요. (OFFICE_MACRO)")
        document = _read_word(name, content) if kind == "docx" else _read_excel(name, content)
        if document.character_count > 3_000_000:
            raise PdfInputError("추출된 글이 너무 많습니다. 문서를 나누어 등록하세요. (TEXT_SIZE)")
        if not document.text_page_count:
            raise PdfInputError("검색할 글이 없습니다. 이미지 문서는 OCR 후 등록하세요. (NO_TEXT)")
        return document
    except PdfInputError:
        raise
    except Exception:
        # Office 파서의 예외에는 셀 값·경로 등이 포함될 수 있습니다.
        raise PdfInputError("문서를 읽지 못했습니다. 암호·파일 형식을 확인하고 다시 저장하세요. (OFFICE_READ)") from None


def _read_word(name, content):
    import re

    from docx import Document
    from docx.table import Table

    document, units, headings = Document(BytesIO(content)), [], {}
    paragraph = table = 0
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            table += 1
            text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in block.rows)
            location = f"표 {table}"
            header = " | ".join(cell.text for cell in block.rows[0].cells) if block.rows else ""
        else:
            paragraph += 1
            text, location = block.text, f"문단 {paragraph}"
            header = ""
            heading = re.search(r"(?:Heading|제목)\s*(\d+)", block.style.name or "", re.I)
            if heading:
                level = int(heading.group(1))
                headings = {k: v for k, v in headings.items() if k < level}
                headings[level] = text.strip()
        if text.strip():
            units.append(PdfPage(None, text.strip(), location=location,
                                 section=" > ".join(headings[k] for k in sorted(headings)), header=header))
    return PdfDocument(name, tuple(units), "docx",
        ("Word의 머리말·각주·텍스트 상자·이미지는 검색에 포함되지 않을 수 있습니다. 문단과 표를 확인하세요.",
         "Word에는 고정 페이지 번호가 없어 문단·표 위치를 표시합니다. 정확한 쪽수가 필요하면 PDF로 등록하세요."))


def _read_excel(name, content):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    values = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    formulas = None
    units = []
    try:
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        for sheet in values:
            if sheet.sheet_state != "visible":
                continue
            if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 100:
                raise PdfInputError("Excel은 시트당 20,000행·100열까지 지원합니다. 필요한 영역만 저장하세요. (EXCEL_SIZE)")
            first_row = None
            for row, formula_row in zip(sheet.iter_rows(), formulas[sheet.title].iter_rows(), strict=True):
                if any(f.data_type == "f" and v.value is None for v, f in zip(row, formula_row, strict=True)):
                    raise PdfInputError("저장된 계산 결과가 없는 수식이 있습니다. Excel에서 다시 계산·저장하거나 값으로 저장하세요. (EXCEL_FORMULA)")
                nonempty = [cell for cell in row if cell.value is not None]
                if not nonempty:
                    continue
                number = nonempty[0].row
                text = " | ".join(f"{c.column_letter}: {c.value}" for c in nonempty)
                span = f"A{number}:{get_column_letter(len(row))}{number}"
                if first_row is None:
                    first_row = (span, text)
                # 열 제목 후보를 함께 보관합니다. 실제 제목 여부는 관리자 미리보기에서 확인합니다.
                prefix = "" if span == first_row[0] else f"첫 데이터 행(열 제목 확인 필요) {first_row[0]}: {first_row[1]}\n"
                units.append(PdfPage(None, prefix + text, location=f"{sheet.title}!{span}",
                                     section=sheet.title))
    finally:
        values.close()
        if formulas is not None:
            formulas.close()
    return PdfDocument(name, tuple(units), "xlsx",
        ("Excel의 보이는 시트와 저장된 셀 값만 검색합니다. 병합 셀·표 제목·이미지의 누락을 확인하세요.",
         "Excel 출처는 인쇄 쪽수 대신 시트·셀 주소로 표시합니다."))
