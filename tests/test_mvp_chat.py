"""가상 원문·가짜 API 응답으로 검증합니다. 실제 환자 자료와 유료 API를 사용하지 않습니다."""

import json
from pathlib import Path

import httpx
import numpy as np
import pytest
from pglast import parse_sql
from streamlit.testing.v1 import AppTest

from mvp.ai import Quota, RateLimitError, estimated_tokens, generate, validate_answer
from mvp.cloud import StaffLibrary
from mvp.documents import PdfDocument, PdfPage
from mvp.library import (
    Chunk,
    Hit,
    LocalLibrary,
    embedding_question,
    make_chunks,
    protect_private,
    rank_hits,
    retrieval_question,
    validate_checklist,
)
from mvp.settings import DIMENSIONS, MODEL, GuideError, Settings, reject_secret_key

APP = Path(__file__).resolve().parents[1] / "mvp" / "app.py"


def part(identifier="chunk-1", document="doc-1", page=3, text="교육실 사용 전 예약 확인표를 확인합니다."):
    return Chunk(identifier, document, "가상교육실.pdf", page, "가상 교육실 안내", "", None, text, 0)


def valid_response(chunk):
    return json.dumps({"answerable": True, "statements": [
        {"text": "교육실 사용 전 예약 확인표를 확인합니다.",
         "evidence": [{"chunk_id": chunk.id, "quote": chunk.text}]}
    ]}, ensure_ascii=False)


class FakeEmbedder:
    def count(self, text):
        return len(text)

    def encode(self, texts):
        vectors = np.zeros((len(texts), DIMENSIONS), dtype=np.float32)
        vectors[:, 0] = 1
        return vectors



@pytest.fixture(autouse=True)
def isolated_repository(monkeypatch, tmp_path):
    import streamlit as st
    monkeypatch.setenv("GUIDE_DATA_DIR", str(tmp_path / "library"))
    monkeypatch.setenv("GUIDE_STORAGE_BACKEND", "local")
    st.cache_resource.clear()


def registered_app(monkeypatch, tmp_path):
    from io import BytesIO

    from docx import Document

    from mvp.auth import LocalAuth
    from mvp.repository import Repository
    from mvp.settings import load_settings
    from mvp.storage import source_store

    settings = load_settings()
    auth = LocalAuth(settings.library_dir)
    auth.bootstrap("admin", "fictional-password")
    library = Repository(settings, auth, source_store(settings, auth))
    word = Document()
    word.add_paragraph("교육실 사용 전 예약 확인표를 확인합니다.")
    word.add_paragraph("사용 후 빔프로젝터의 전원을 끕니다.")
    buffer = BytesIO()
    word.save(buffer)
    doc_id = library.register("가상교육실.docx", buffer.getvalue(), FakeEmbedder())
    parts = library.source_chunks(doc_id)
    items = [{"chunk_id": c.id, "quote": c.text, "stage": "사전 확인"} for c in parts]
    library.save_checklist(validate_checklist("교육실 이용", ["교육실"], items, parts))
    app = AppTest.from_file(str(APP), default_timeout=30).run()
    app.session_state["auth"] = auth
    app.run()
    return app


def ask(app, question):
    app.chat_input[0].set_value(question).run()


def test_chunks_keep_blank_page_locations_and_unknown_metadata():
    document = PdfDocument("교육실.pdf", (PdfPage(1, "안내"), PdfPage(2, ""), PdfPage(3, "예약을 확인합니다." * 30)))
    metadata, chunks = make_chunks(document, FakeEmbedder(), file_hash="a" * 64)
    assert {c.page for c in chunks} == {1, 3}
    assert all(c.updated_date is None and c.section == "" for c in chunks)
    assert all(c.document_id == metadata["id"] and len(c.text) <= 110 for c in chunks)
    assert "예약을 확인합니다." in chunks[-1].text


def test_vector_search_scopes_documents_and_preserves_english_abbreviations():
    library = LocalLibrary()
    chunks = [part(text="VRE 확인 지침"), part("chunk-2", "doc-2", text="CRE 확인 지침")]
    library.docs = [{"id": "doc-1"}, {"id": "doc-2"}]
    library.chunks = chunks
    library.vectors = FakeEmbedder().encode(["a", "b"])
    vector = FakeEmbedder().encode(["q"])[0]
    assert library.search("VRE", vector, ["doc-2"], 0.1) == []
    assert library.search("VRE", vector, ["doc-1", "doc-2"], 0.1)[0].chunk.id == "chunk-1"
    assert {h.chunk.id for h in library.search("VRE CRE 교육자료 비교", vector, ["doc-1", "doc-2"], 0.1)} == {
        "chunk-1", "chunk-2"
    }
    assert rank_hits("PCN", [Hit(chunks[0], 0.99)]) == []
    assert rank_hits("주차장", [Hit(chunks[0], 0.1)]) == []


def test_followup_is_explicit_and_does_not_use_previous_ai_answer_as_evidence():
    assert retrieval_question("준비물은?", "교육실 이용", True) == "교육실 이용 / 추가 질문: 준비물은?"
    assert retrieval_question("다른 질문", "교육실 이용", False) == "다른 질문"


def test_korean_questions_retrieve_exact_source_despite_low_embedding_score():
    # 임상 지시가 없는 가상 문서로 사용자가 입력한 표현을 재현합니다.
    topic = Chunk("purpose", "sedation", "진정간호_가상자료.pdf", 1, "진정간호 교육", "", None,
                  "진정 간호\n1. 목적\n이 문서는 검색 기능을 시험하는 가상 자료입니다.", 0)
    procedure = Chunk("procedure", "sedation", topic.document_name, 2, topic.title, "", None,
                      "6. 진정 절차\n이 문단은 검색 시험을 위한 가상 문단입니다.", 1)
    unrelated = part("other", text="교육실 사용 시간 안내", page=9)
    candidates = [Hit(unrelated, .50), Hit(topic, .28), Hit(procedure, .17)]
    for query in ("진정간호란", "진정간호 목적에 대해 알려줘", "진정 간호 목적은?"):
        result = rank_hits(query, candidates)
        assert result[0].chunk.id == "purpose"
        assert result[0].chunk.text == topic.text  # 검색 정규화가 출처 원문을 바꾸지 않음
    assert "procedure" in {h.chunk.id for h in rank_hits("진정간호 절차에 대해 알려줘", candidates)}
    assert rank_hits("퇴직금 계산 방법", candidates) == []
    assert embedding_question("진정간호 목적에 대해 알려줘") == "진정 간호 목적"
    assert embedding_question("진정간호란") == "진정 간호"


def test_document_title_alone_and_generic_intent_are_not_keyword_evidence():
    chunk = Chunk("x", "d", "진정간호.pdf", 3, "진정간호", "", None, "자료 배포일 안내", 1)
    assert rank_hits("진정간호 투여량은?", [Hit(chunk, .20)]) == []
    assert rank_hits("PCN 목적", [Hit(part(text="교육실 목적 안내"), .99)]) == []
    assert rank_hits("퇴직금 목적", [Hit(part(text="교육실 목적 안내"), .45)]) == []


def test_local_search_includes_keyword_sources_outside_dense_top_forty():
    library = LocalLibrary()
    library.chunks = [part(str(i), text="검색과 관계없는 가상 문단") for i in range(45)]
    source = part("target", page=2, text="진정 간호\n목적\n가상 교육자료 검색 안내입니다.")
    library.chunks.append(source)
    library.vectors = FakeEmbedder().encode(["x"] * 46)
    library.vectors[-1, 0] = .20
    library.vectors[-1, 1] = np.sqrt(1 - .20 ** 2)
    vector = FakeEmbedder().encode(["q"])[0]
    hits = library.search("진정간호 목적에 대해 알려줘", vector, ["doc-1"], .38)
    assert hits[0].chunk.id == source.id


@pytest.mark.parametrize("text", ["환자명: 홍길동", "MRN: 123456", "환자등록번호: 1234567", "010-1234-5678"])
def test_private_identifiers_blocked_before_search_or_network(text):
    with pytest.raises(GuideError, match="PRIVACY"):
        protect_private(text)


def test_answer_requires_exact_real_citations_and_no_invented_numbers():
    chunk = part()
    hit = Hit(chunk, 0.8)
    assert validate_answer(valid_response(chunk), [hit]).answerable
    for raw in [
        valid_response(chunk).replace("chunk-1", "made-up-source"),
        valid_response(chunk).replace(chunk.text, "없는 원문 문장입니다."),
        json.dumps({"answerable": True, "statements": [{"text": "3일 전에 예약합니다.",
            "evidence": [{"chunk_id": chunk.id, "quote": chunk.text}]}]}),
        '{"answerable":true,"statements":[]}',
        '{"answerable":false,"statements":[],"extra":"invented"}',
    ]:
        with pytest.raises(GuideError, match="AI_EVIDENCE"):
            validate_answer(raw, [hit])
    assert not validate_answer('{"answerable":false,"statements":[]}', [hit]).answerable


def test_approved_checklist_requires_original_quotes_and_single_document():
    c1, c2 = part(), part("chunk-2", "doc-2")
    items = [{"chunk_id": c1.id, "quote": c1.text, "stage": "사전 확인"}]
    assert validate_checklist("교육실", ["교육실"], items, [c1])["document_id"] == c1.document_id
    with pytest.raises(GuideError):
        validate_checklist("교육실", ["교육실"], [{**items[0], "quote": "원문에 없는 항목"}], [c1])
    with pytest.raises(GuideError):
        validate_checklist("교육실", ["교육실"], items + [{**items[0], "chunk_id": c2.id}], [c1, c2])


def test_quota_is_persistent_shared_and_reserves_failed_requests(tmp_path):
    settings = Settings(daily_limit=2, user_daily_limit=1)
    path = tmp_path / "quota.sqlite3"
    Quota(path).reserve(settings, "employee-a", 100, now=1000)
    with pytest.raises(GuideError, match="AI_LIMIT"):
        Quota(path).reserve(settings, "employee-a", 100, now=1001)
    Quota(path).reserve(settings, "employee-b", 100, now=1001)
    with pytest.raises(GuideError, match="AI_LIMIT"):
        Quota(path).reserve(settings, "employee-c", 100, now=1001)
    Quota(path).reserve(settings, "employee-a", 100, now=100000)


def test_free_token_limit_stops_request_without_wait_or_paid_fallback(tmp_path):
    settings = Settings(llm_provider="groq_free")
    quota = Quota(tmp_path / "quota.sqlite3")
    quota.reserve(settings, "a", 6500, now=1000)
    with pytest.raises(GuideError, match="AI_LIMIT"):
        quota.reserve(settings, "b", 6500, now=1001)


def test_llm_disabled_does_not_call_network_and_empty_search_abstains():
    def fail(request):
        pytest.fail("No API calls are allowed")
    settings = Settings()
    transport = httpx.MockTransport(fail)
    with pytest.raises(GuideError, match="AI_SETUP"):
        generate(settings, "교육실", [Hit(part(), 0.8)], "a", transport=transport)
    answer, _ = generate(settings, "교육실", [], "a", transport=transport)
    assert not answer.answerable


def test_single_ai_call_uses_only_selected_chunks_and_returns_validated_answer(tmp_path):
    calls = []
    chunk = part()
    def handle(request):
        calls.append(request)
        data = json.loads(request.content)
        assert data["model"] == "hospital-model"
        assert "tools" not in data and "stream" not in data
        assert chunk.text in data["messages"][1]["content"]
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop",
            "message": {"content": valid_response(chunk)}}]})
    settings = Settings(llm_provider="internal", llm_url="https://example.invalid/v1",
                        llm_model="hospital-model", llm_approved=True)
    answer, sources = generate(settings, "교육실 사용 전 확인 사항", [Hit(chunk, 0.8)], "a",
                               Quota(tmp_path / "quota.sqlite3"), httpx.MockTransport(handle))
    assert answer.answerable and sources[0].chunk.id == chunk.id and len(calls) == 1


@pytest.mark.parametrize("status,code", [(429,"AI_RATE"), (401,"AI_AUTH"), (500,"AI_SERVER")])
def test_api_failure_is_not_misreported_as_missing_guideline(tmp_path, status, code):
    settings = Settings(llm_provider="internal", llm_url="https://example.invalid/v1",
                        llm_model="hospital-model", llm_approved=True)
    with pytest.raises(GuideError, match=code):
        generate(settings, "교육실", [Hit(part(), .8)], "a", Quota(tmp_path / "quota.sqlite3"),
                 httpx.MockTransport(lambda r: httpx.Response(status)))


def test_auth_rechecks_membership_and_drops_session_when_employee_revoked():
    active = True
    headers = []
    def handle(request):
        headers.append(dict(request.headers))
        if request.url.path == "/auth/v1/token":
            return httpx.Response(200, json={"access_token": "user-jwt", "refresh_token": "refresh",
                                            "expires_in": 3600})
        if request.url.path == "/auth/v1/user":
            return httpx.Response(200, json={"id": "employee-1"})
        if request.url.path == "/rest/v1/guide_profiles":
            return httpx.Response(200, json=[{"user_id": "employee-1", "role": "staff", "active": active}])
        pytest.fail("Unexpected endpoint")
    library = StaffLibrary(Settings(supabase_url="https://example.invalid", supabase_key="public-key"),
                           httpx.MockTransport(handle))
    assert library.login("fictional@example.invalid", "test-only")["role"] == "staff"
    assert headers[-1]["authorization"] == "Bearer user-jwt"
    with pytest.raises(GuideError, match="관리자"):
        library.require_admin()
    active = False
    with pytest.raises(GuideError, match="ACCESS"):
        library.authorize()
    assert not library.token and library.profile is None


def test_secret_key_and_non_https_ai_endpoint_rejected():
    import base64
    encoded = base64.urlsafe_b64encode(b'{"role":"service_role"}').decode().rstrip("=")
    for key in ("sb_secret_not-allowed", "header." + encoded + ".signature"):
        with pytest.raises(GuideError):
            reject_secret_key(key)
    with pytest.raises(GuideError, match="AI_URL"):
        Settings(llm_provider="internal", llm_url="http://remote.invalid/v1",
                 llm_model="x", llm_approved=True).llm_endpoint()


def test_sql_is_parseable_and_limits_writes_and_anonymous_access():
    sql = (APP.parent / "schema.sql").read_text(encoding="utf-8")
    assert len(parse_sql(sql)) >= 35
    assert "security invoker" in sql
    assert "ADMIN_REQUIRED" in sql and "INVALID_EVIDENCE" in sql
    assert "to anon" not in sql.lower()
    assert "grant update(status)" in sql
    assert "grant insert" not in sql.lower()
    assert MODEL in sql


def test_staff_without_configuration_fails_closed(monkeypatch):
    monkeypatch.setenv("GUIDE_MODE", "staff")
    monkeypatch.setenv("GUIDE_SUPABASE_URL", "")
    monkeypatch.setenv("GUIDE_SUPABASE_PUBLISHABLE_KEY", "")
    app = AppTest.from_file(str(APP), default_timeout=30).run()
    assert not app.exception
    assert any("로그인 설정" in x.value for x in app.subheader)
    assert not app.chat_input


def test_chat_form_sources_and_new_conversation_without_checklist_ui(monkeypatch, tmp_path):
    monkeypatch.setenv("GUIDE_MODE", "local")
    monkeypatch.setenv("GUIDE_LLM_PROVIDER", "disabled")
    monkeypatch.setattr("mvp.library.Embedder", FakeEmbedder)
    app = registered_app(monkeypatch, tmp_path)
    assert not app.exception
    assert app.chat_input[0].disabled is False
    ask(app, "교육실은 어떻게 사용해?")
    assert not app.exception
    assert any("AI_SETUP" in x.value for x in app.info)
    assert not any("체크리스트" in b.label for b in app.button)
    assert any("문단 2" in e.label for e in app.expander)
    assert all("체크리스트" not in e.label and "PDF 원문 확인" not in e.label for e in app.expander)
    next(b for b in app.button if b.label == "새 대화 시작").click().run()
    assert not app.exception
    assert app.session_state["turns"] == []
    assert app.chat_input[0].value in (None, "")
    assert next(c for c in app.checkbox if c.label == "이전 질문에 이어서 묻기").disabled


def test_ai_answer_is_rendered_in_chat_after_real_validation(monkeypatch, tmp_path):
    monkeypatch.setenv("GUIDE_MODE", "local")
    monkeypatch.setenv("GUIDE_LLM_PROVIDER", "internal")
    monkeypatch.setenv("GUIDE_LLM_BASE_URL", "https://example.invalid/v1")
    monkeypatch.setenv("GUIDE_LLM_MODEL", "hospital-model")
    monkeypatch.setenv("GUIDE_LLM_APPROVED", "true")
    monkeypatch.setattr("mvp.library.Embedder", FakeEmbedder)
    calls = []
    def response(request):
        calls.append(request)
        content = json.loads(request.content)["messages"][1]["content"]
        evidence = json.loads(content.split("Evidence (JSON):\n", 1)[1])
        row = next(e for e in evidence if "예약 확인표" in e["text"])
        data = {"answerable": True, "statements": [{"text": "교육실 사용 전 예약 확인표를 확인합니다.",
            "evidence": [{"chunk_id": row["chunk_id"], "quote": "교육실 사용 전 예약 확인표를 확인합니다."}]}]}
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop",
            "message": {"content": json.dumps(data, ensure_ascii=False)}}]})
    def invoke(settings, question, hits, user_id, **kwargs):
        return generate(settings, question, hits, user_id, Quota(tmp_path / "ui.sqlite3"),
                        httpx.MockTransport(response), **kwargs)
    monkeypatch.setattr("mvp.ai.generate", invoke)
    app = registered_app(monkeypatch, tmp_path)
    ask(app, "교육실 사용 전 확인할 사항은?")
    assert not app.exception
    assert len(calls) == 1
    assert any("교육실 사용 전 예약 확인표를 확인합니다" in t.value and "[[1]](#source-" in t.value for t in app.markdown)
    assert not any("관련 원문" in e.label for e in app.expander)
    next(b for b in app.button if b.key and b.key.startswith("open_source_")).click().run()
    assert len(calls) == 1
    assert any("문단 1" in e.label for e in app.expander)
    assert not any("AI_SETUP" in t.value for t in app.info)
    assert any(t.value == "근거 1" for t in app.caption)
    assert any("관련 원문 펼쳐보기" in e.label for e in app.expander)


def test_typing_and_empty_submission_do_not_call_ai(monkeypatch, tmp_path):
    monkeypatch.setenv("GUIDE_MODE", "local")
    monkeypatch.setattr("mvp.library.Embedder", FakeEmbedder)
    monkeypatch.setattr("mvp.ai.generate", lambda *args: pytest.fail("Unsubmitted input must not call AI"))
    app = registered_app(monkeypatch, tmp_path)
    app.session_state["question_input"] = "작성 중인 질문"
    app.run()
    assert app.session_state["turns"] == []
    ask(app, "   ")
    assert not app.exception and app.session_state["turns"] == []
    assert any("질문을 입력" in w.value for w in app.warning)


def test_follow_up_form_and_new_conversation_use_only_question_context(monkeypatch, tmp_path):
    monkeypatch.setenv("GUIDE_MODE", "local")
    monkeypatch.setattr("mvp.library.Embedder", FakeEmbedder)
    calls = []
    def answer(settings, question, hits, user_id, **kwargs):
        calls.append(question)
        return validate_answer('{"answerable":false,"statements":[]}', hits), hits
    monkeypatch.setattr("mvp.ai.generate", answer)
    app = registered_app(monkeypatch, tmp_path)
    ask(app, "교육실 예약 확인")
    next(c for c in app.checkbox if c.label == "이전 질문에 이어서 묻기").check().run()
    ask(app, "사용 후에는?")
    assert not app.exception
    assert calls == ["교육실 예약 확인", "교육실 예약 확인 / 추가 질문: 사용 후에는?"]
    assert any(c.value == "이전 질문의 맥락을 이어서 확인했습니다." for c in app.caption)
    assert any(m.value == "등록된 지침서에서 확인할 수 없습니다." for m in app.markdown)
    next(b for b in app.button if b.label == "새 대화 시작").click().run()
    ask(app, "교육실 사용 후 정리")
    assert calls[-1] == "교육실 사용 후 정리"
    assert len(app.session_state["turns"]) == 1


def test_groq_receives_only_locally_retrieved_chunks_and_citations_are_grouped(monkeypatch, tmp_path):
    from io import BytesIO

    from docx import Document

    from mvp.auth import LocalAuth
    from mvp.repository import Repository
    from mvp.settings import load_settings
    from mvp.storage import source_store

    monkeypatch.setenv("GUIDE_MODE", "local")
    monkeypatch.setenv("GUIDE_LLM_PROVIDER", "groq_free")
    monkeypatch.setenv("GUIDE_LLM_MODEL", "openai/gpt-oss-20b")
    monkeypatch.setenv("GUIDE_LLM_API_KEY", "synthetic-key")
    monkeypatch.setenv("GUIDE_LLM_APPROVED", "true")
    monkeypatch.setenv("GUIDE_GROQ_FREE_CONFIRMED", "true")
    monkeypatch.setattr("mvp.library.Embedder", FakeEmbedder)
    app = registered_app(monkeypatch, tmp_path)
    settings, auth = load_settings(), app.session_state["auth"]
    repo = Repository(settings, auth, source_store(settings, auth))
    word = Document()
    source = "VRE 검색 시험용 가상 안내문입니다. 임상 지시를 포함하지 않습니다."
    excluded = "이 별도 문단은 Groq로 보내지 않아야 합니다."
    word.add_paragraph(source)
    word.add_paragraph(excluded)
    buffer = BytesIO()
    word.save(buffer)
    repo.register("가상검색검증.docx", buffer.getvalue(), FakeEmbedder())
    auth.create_user("reader", "synthetic-reader-password")
    staff = LocalAuth(settings.library_dir)
    staff.login("reader", "synthetic-reader-password")
    app.session_state["auth"] = staff
    captured = []
    def response(request):
        captured.append(request)
        assert str(request.url) == "https://api.groq.com/openai/v1/chat/completions"
        payload = json.loads(request.content)
        content = payload["messages"][1]["content"]
        assert excluded not in content and "교육실" not in content
        evidence = json.loads(content.split("Evidence (JSON):\n", 1)[1])
        assert len(evidence) == 1 and evidence[0]["text"] == source
        item = {"text": source, "evidence": [{"chunk_id": evidence[0]["chunk_id"], "quote": source}]}
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {
            "content": json.dumps({"answerable": True, "statements": [item, item]}, ensure_ascii=False)}}]})
    def invoke(config, question, hits, user_id, **kwargs):
        return generate(config, question, hits, user_id, Quota(tmp_path / "groq.sqlite3"),
                        httpx.MockTransport(response), **kwargs)
    monkeypatch.setattr("mvp.ai.generate", invoke)
    # 질문 시 원본 저장소를 다시 읽거나 체크리스트를 조회하면 실패하도록 검증합니다.
    monkeypatch.setattr("mvp.storage.LocalSourceStore.read", lambda *a: pytest.fail("Original read during chat"))
    monkeypatch.setattr("mvp.repository.Repository.list_checklists", lambda *a: pytest.fail("Checklist during chat"))
    app.run()
    ask(app, "VRE 검색 시험")
    assert not app.exception and len(captured) == 1
    assert not app.tabs and not app.get("file_uploader")
    assert not any("관련 원문" in e.label for e in app.expander)
    next(b for b in app.button if b.key and b.key.startswith("open_source_")).click().run()
    assert len(captured) == 1
    assert len([e for e in app.expander if "관련 원문 펼쳐보기" in e.label]) == 1
    assert any(t.value == "가상검색검증.docx · 문단 1" for t in app.text)


def test_concurrent_quota_cannot_exceed_global_limit(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    quota = Quota(tmp_path / "concurrent.sqlite3")
    settings = Settings(daily_limit=2, user_daily_limit=2)
    def reserve(index):
        try:
            quota.reserve(settings, str(index), 100, now=1000)
            return True
        except GuideError:
            return False
    with ThreadPoolExecutor(max_workers=4) as workers:
        accepted = list(workers.map(reserve, range(4)))
    assert sum(accepted) == 2


def test_korean_input_uses_model_tokens_instead_of_utf8_byte_count():
    from mvp.ai import OUTPUT_LIMIT
    messages = [{"role": "user", "content": "가상 교육실 안내에서 예약 확인 방법을 알려주세요. " * 30}]
    byte_reservation = len(json.dumps(messages, ensure_ascii=False).encode()) + OUTPUT_LIMIT
    estimate = estimated_tokens(messages, "groq_free")
    assert OUTPUT_LIMIT < estimate < byte_reservation * .75
    assert estimated_tokens(messages, "internal") == byte_reservation


def test_three_completed_questions_settle_actual_usage_and_keep_call_counts(tmp_path):
    import sqlite3
    settings = Settings(llm_provider="groq_free", daily_limit=10, user_daily_limit=10)
    quota = Quota(tmp_path / "actual.sqlite3")
    for second in range(3):
        identifier = quota.reserve(settings, "a", 4000, now=1000 + second)
        quota.settle(identifier, {"total_tokens": 1000}, now=1000 + second)
    with sqlite3.connect(quota.path) as db:
        assert db.execute("select count(*),sum(tokens),sum(reported) from reservations").fetchone() == (3,3000,3)
    # 완료된 요청의 실제 사용량이 있어도 동시에 실행 중인 최대 예약량은 보호합니다.
    quota.reserve(settings, "b", 4000, now=1003)
    with pytest.raises(RateLimitError, match="AI_LIMIT_MINUTE") as error:
        quota.reserve(settings, "c", 2000, now=1003)
    assert error.value.retry_after == 57


def test_settlement_is_idempotent_and_accepts_actual_usage_above_estimate(tmp_path):
    import sqlite3
    quota = Quota(tmp_path / "once.sqlite3")
    identifier = quota.reserve(Settings(), "a", 100, now=1000)
    quota.settle(identifier, {"total_tokens": 200}, now=1001)
    quota.settle(identifier, {"total_tokens": 1}, now=1002)
    with sqlite3.connect(quota.path) as db:
        assert db.execute("select tokens,reported from reservations").fetchone() == (200,1)


@pytest.mark.parametrize("usage", [None, {}, {"total_tokens": "10"}, {"total_tokens": True},
                                 {"total_tokens": -1}, {"total_tokens": 0}])
def test_unknown_usage_keeps_reservation(tmp_path, usage):
    import sqlite3
    quota = Quota(tmp_path / "unknown.sqlite3")
    identifier = quota.reserve(Settings(), "a", 4000, now=1000)
    quota.settle(identifier, usage, now=1001)
    with sqlite3.connect(quota.path) as db:
        assert db.execute("select tokens,reported from reservations").fetchone() == (4000,0)


def test_minute_wait_expires_and_daily_limit_has_distinct_message(tmp_path):
    quota = Quota(tmp_path / "wait.sqlite3")
    settings = Settings(llm_provider="groq_free")
    identifier = quota.reserve(settings, "a", 6500, now=1000)
    with pytest.raises(RateLimitError) as error:
        quota.reserve(settings, "a", 2000, now=1010)
    assert error.value.retry_after == 50
    assert "하루 사용량 소진은 아닙니다" in str(error.value)
    quota.reserve(settings, "a", 2000, now=1060)
    quota.settle(identifier, {"total_tokens": 199000}, now=1061)
    with pytest.raises(GuideError, match="AI_LIMIT_DAY"):
        quota.reserve(settings, "a", 2000, now=1122)


def test_existing_quota_records_migrate_without_reset(tmp_path):
    import sqlite3
    path = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(path) as db:
        db.execute("create table reservations(at real,user_hash text,tokens integer)")
        db.execute("insert into reservations values (1000,'old-user-hash',4352)")
    Quota(path)
    Quota(path)
    with sqlite3.connect(path) as db:
        assert db.execute("select at,tokens,reported from reservations").fetchall() == [(1000,4352,0)]


def test_successful_api_usage_settles_even_if_answer_evidence_is_invalid(tmp_path):
    import sqlite3
    quota = Quota(tmp_path / "invalid-answer.sqlite3")
    settings = Settings(llm_provider="groq_free", llm_model="openai/gpt-oss-20b",
                        llm_key="fake-key", groq_free_confirmed=True, llm_approved=True)
    response = {"choices": [{"finish_reason": "stop", "message": {"content": "invalid-json"}}],
                "usage": {"prompt_tokens": 200, "completion_tokens": 100, "total_tokens": 300}}
    with pytest.raises(GuideError, match="AI_EVIDENCE"):
        generate(settings, "교육실", [Hit(part(),.8)], "a", quota,
                 httpx.MockTransport(lambda request: httpx.Response(200,json=response)))
    with sqlite3.connect(quota.path) as db:
        assert db.execute("select tokens,reported from reservations").fetchone() == (300,1)


def test_provider_retry_after_is_shown_without_automatic_request(tmp_path):
    calls = []
    def handle(request):
        calls.append(request)
        return httpx.Response(429, headers={"retry-after": "12.5"})
    settings = Settings(llm_provider="groq_free", llm_model="openai/gpt-oss-20b",
                        llm_key="fake-key", groq_free_confirmed=True, llm_approved=True)
    with pytest.raises(RateLimitError, match="AI_RATE") as error:
        generate(settings, "교육실", [Hit(part(),.8)], "a", Quota(tmp_path / "429.sqlite3"),
                 httpx.MockTransport(handle))
    assert error.value.retry_after == 13
    assert len(calls) == 1


def test_retry_button_waits_and_reuses_original_question_context(monkeypatch, tmp_path):
    import time
    monkeypatch.setenv("GUIDE_MODE", "local")
    monkeypatch.setenv("GUIDE_LLM_PROVIDER", "disabled")
    monkeypatch.setattr("mvp.library.Embedder", FakeEmbedder)
    calls = []
    def limited_then_answer(settings, question, hits, user_id, **kwargs):
        calls.append(question)
        if len(calls) == 1:
            raise RateLimitError("분당 대기 안내 (AI_LIMIT_MINUTE)", 60)
        chunk = hits[0].chunk
        response = json.dumps({"answerable": True, "statements": [
            {"text": chunk.text, "evidence": [{"chunk_id": chunk.id, "quote": chunk.text}]}]})
        return validate_answer(response, hits), hits
    monkeypatch.setattr("mvp.ai.generate", limited_then_answer)
    app = registered_app(monkeypatch, tmp_path)
    ask(app, "교육실 예약 확인")
    assert not app.exception
    next(b for b in app.button if b.label == "같은 질문 다시 시도").click().run()
    assert len(calls) == 1
    assert any("남았습니다" in w.value for w in app.warning)
    next(c for c in app.checkbox if c.label == "이전 질문에 이어서 묻기").check().run()
    app.session_state["turns"][-1]["retry_at"] = time.time() - 1
    next(b for b in app.button if b.label == "같은 질문 다시 시도").click().run()
    assert not app.exception
    assert calls == ["교육실 예약 확인", "교육실 예약 확인"]
    assert app.session_state["turns"][-1]["answer"].answerable


def test_identical_question_reuses_only_current_session_validated_answer(monkeypatch, tmp_path):
    monkeypatch.setenv('GUIDE_MODE', 'local')
    monkeypatch.setattr('mvp.library.Embedder', FakeEmbedder)
    calls = []
    def respond(settings, question, hits, user_id, **kwargs):
        calls.append(question)
        source = hits[0].chunk
        return validate_answer(json.dumps({'answerable': True, 'statements': [{'text': source.text,
            'evidence': [{'chunk_id': source.id, 'quote': source.text}]}]}), hits), hits
    monkeypatch.setattr('mvp.ai.generate', respond)
    app = registered_app(monkeypatch, tmp_path)
    ask(app, '교육실 예약 확인')
    ask(app, '교육실 예약 확인')
    assert not app.exception and len(calls) == 1
    assert app.session_state['turns'][-1]['reused']
    next(b for b in app.button if b.label == '새 대화 시작').click().run()
    ask(app, '교육실 예약 확인')
    assert len(calls) == 2
