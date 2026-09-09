"""문맥 누락·단위 오류·등록 상태를 가상 자료로 검사합니다. 외부 AI 호출 없음."""

import json
from dataclasses import replace
from io import BytesIO

import numpy as np
import pytest
from docx import Document
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from mvp.ai import Quota, validate_answer
from mvp.context import neighbors
from mvp.documents import read_document
from mvp.library import (
    Chunk,
    Hit,
    LocalLibrary,
    anchors,
    bounded_embedding_question,
    make_chunks,
    retrieval_question,
)
from mvp.settings import DIMENSIONS, GuideError, Settings


class Model:
    def count(self, text):
        return len(text)


def chunk(identifier, text, index=0, page=1, section="PCN 교육"):
    return Chunk(identifier, "synthetic", "가상교육.pdf", page, "가상교육", section, None, text, index)


def test_continuation_retrieved_without_repeating_acronym_and_other_section_excluded():
    parts = [chunk("a", "PCN 가상 교육 절차", 0),
             chunk("b", "주의사항: 교육표를 생략하지 않습니다.", 1, page=2),
             chunk("c", "별도 항목의 설명", 2, page=2, section="다른 주제")]
    library = LocalLibrary()
    library.chunks = parts
    library.vectors = np.zeros((3, DIMENSIONS), dtype=np.float32)
    library.vectors[0, 0] = 1
    query = library.vectors[0]
    hits = library.search("PCN 절차", query, ["synthetic"], .38)
    assert {h.chunk.id for h in hits} == {"a", "b"}
    assert next(h for h in hits if h.chunk.id == "b").chunk.page == 2


def test_context_stops_at_competing_acronym_and_document_boundaries():
    a = chunk("a", "VRE 교육", section="")
    b = chunk("b", "CRE 교육", index=1, section="")
    c = replace(chunk("c", "추가 설명", index=1, section=""), document_id="other")
    assert neighbors(a, [a, b, c]) == [a]


def test_requested_purpose_ranks_above_generic_high_vector_score():
    from mvp.library import rank_hits
    generic = chunk('general', 'PCN 가상 교육 자료를 배포합니다.')
    purpose = chunk('purpose', 'PCN 교육 목적은 검색 기능 검증입니다.', index=1)
    hits = rank_hits('PCN 목적', [Hit(generic, .99), Hit(purpose, .25)])
    assert hits[0].chunk.id == 'purpose'


def test_korean_aliases_do_not_confuse_drug_and_resistance_types():
    assert anchors("경피적신루 관리") == ["pcn"]
    assert anchors("말초삽입중심정맥관") == ["picc"]
    assert anchors("반코마이신 내성 장알균") == ["vre"]
    assert set(anchors("CRE CPE 비교")) == {"cre", "cpe"}
    assert anchors("PCNA") == []


def test_followup_keeps_recent_context_and_resets_when_acronym_changes():
    context = "PCN 교육"
    for index in range(15):
        context = retrieval_question(f"추가 조건 {index}는?", context, True)
    assert "PCN 교육" in context and "조건 14" in context and "조건 13" in context
    assert len(context) < 500
    assert retrieval_question("Foley 교육", context, True) == "Foley 교육"
    assert len(bounded_embedding_question("PCN " + "교육표 " * 300, Model())) <= 128


def test_same_number_with_changed_unit_is_rejected():
    source = chunk("a", "가상 단위 검증: 5 mg이라고 표기합니다.")
    payload = {"answerable": True, "statements": [{"text": "5 mL이라고 표기합니다.",
               "evidence": [{"chunk_id": "a", "quote": source.text}]}]}
    with pytest.raises(GuideError, match="AI_EVIDENCE"):
        validate_answer(json.dumps(payload), [Hit(source, .9)])
    payload["statements"][0]["text"] = source.text
    assert validate_answer(json.dumps(payload), [Hit(source, .9)]).answerable


def test_word_table_chunks_retain_header_and_location():
    doc = Document()
    doc.add_heading("교육실 물품", 1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "물품"
    table.rows[0].cells[1].text = "위치"
    for i in range(20):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = f"교육표{i}", "가상교육실"
    buffer = BytesIO()
    doc.save(buffer)
    metadata, parts = make_chunks(read_document("table.docx", buffer.getvalue()), Model())
    table_parts = [p for p in parts if p.location == "표 1"]
    assert len(table_parts) > 2 and all("물품 | 위치" in p.text for p in table_parts)
    assert all(p.page is None and len(p.text) <= 110 for p in parts)
    assert metadata["extraction_version"] == 2


def test_pdf_table_retains_row_relationship_and_real_page():
    buffer = BytesIO()
    table = Table([["Item", "Location"], ["Training card", "Room A"], ["Marker", "Room B"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.black)]))
    SimpleDocTemplate(buffer).build([table])
    doc = read_document("table.pdf", buffer.getvalue())
    assert "Training card | Room A" in doc.pages[0].text
    assert doc.pages[0].number == 1 and doc.warnings


def test_quota_summary_expires_old_records_and_contains_no_user_data(tmp_path):
    quota = Quota(tmp_path / "usage.sqlite3")
    settings = Settings(daily_limit=40)
    quota.reserve(settings, "private-user", 100, now=0)
    quota.reserve(settings, "private-user", 200, now=90000)
    summary = quota.summary(settings, now=90001)
    assert summary["calls"] == 1 and summary["calls_remaining"] == 39 and summary["tokens"] == 200
    assert "private-user" not in json.dumps(summary)


def test_real_local_ocr_reads_synthetic_korean_scan_and_preserves_page():
    from pathlib import Path

    from PIL import Image, ImageDraw, ImageFont
    from pypdf import PdfReader, PdfWriter

    from mvp.pdf_layout import ocr_status
    if not ocr_status()[0]:
        pytest.skip('Optional Tesseract kor+eng runtime is not installed')
    font_path = Path('C:/Windows/Fonts/malgun.ttf')
    if not font_path.is_file():
        pytest.skip('Windows Korean font is not available')
    image = Image.new('RGB', (1600, 800), 'white')
    draw = ImageDraw.Draw(image)
    draw.text((100, 120), 'PCN 교육실 안내', font=ImageFont.truetype(str(font_path), 70), fill='black')
    buffer = BytesIO()
    image.save(buffer, format='PDF', resolution=150)
    writer = PdfWriter()
    writer.add_blank_page(600, 800)
    writer.add_page(PdfReader(buffer).pages[0])
    combined = BytesIO()
    writer.write(combined)
    document = read_document('scan.pdf', combined.getvalue(), ocr=True)
    assert not document.pages[0].text
    assert document.pages[1].number == 2 and 'PCN' in document.pages[1].text
    assert '교육실' in document.pages[1].text
    assert any('OCR로 읽은' in warning for warning in document.warnings)


def test_evaluation_reports_expected_page_coverage_and_latency_without_questions():
    from mvp.evaluate import evaluate
    library = LocalLibrary()
    library.docs = [{'id': 'synthetic'}]
    library.chunks = [chunk('a', 'PCN 가상 교육 절차')]
    library.vectors = np.zeros((1, DIMENSIONS), dtype=np.float32)
    library.vectors[0, 0] = 1
    class VectorModel(Model):
        def encode(self, texts):
            return library.vectors
    report = evaluate(library, VectorModel(), [{'question': 'PCN 교육', 'expected_sources': [
        {'document_name': '가상교육.pdf', 'page': 1}]}])
    assert report['passed'] == 1 and report['evidence_recall'] == 1
    assert report['median_seconds'] >= 0
    assert 'PCN' not in json.dumps(report)
