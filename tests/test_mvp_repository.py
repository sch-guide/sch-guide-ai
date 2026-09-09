"""권한·지속 저장·문서 교체·출처를 합성 문서로 검사합니다. 실제 API는 호출하지 않습니다."""

import json
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import httpx
import numpy as np
import pytest
from docx import Document
from openpyxl import Workbook
from pglast import parse_sql
from reportlab.pdfgen import canvas
from streamlit.testing.v1 import AppTest

from mvp.auth import LocalAuth
from mvp.cloud import StaffLibrary
from mvp.documents import PdfInputError, read_document
from mvp.library import NO_GUIDELINE, validate_checklist
from mvp.repository import Repository, database
from mvp.settings import DIMENSIONS, GuideError, Settings
from mvp.storage import LocalSourceStore, SupabaseSourceStore, source_store

APP = Path(__file__).resolve().parents[1] / "mvp" / "app.py"


def ask(app, question):
    app.chat_input[0].set_value(question).run()


class Model:
    def count(self, text):
        return len(text)

    def encode(self, texts):
        result = np.zeros((len(texts), DIMENSIONS), dtype=np.float32)
        result[:, 0] = 1
        return result


def word_bytes(text="교육실 사용 전 예약 확인표를 확인합니다."):
    doc = Document()
    doc.add_heading("교육실 안내", level=1)
    doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "장소"
    table.cell(0, 1).text = "교육실"
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def setup(tmp_path):
    settings = Settings(mode="local", data_dir=str(tmp_path / "library"))
    admin = LocalAuth(settings.library_dir)
    admin.bootstrap("admin", "synthetic-admin-password")
    admin.create_user("staff", "synthetic-staff-password")
    staff = LocalAuth(settings.library_dir)
    staff.login("staff", "synthetic-staff-password")
    repo = Repository(settings, admin, source_store(settings, admin))
    employee = Repository(settings, staff, source_store(settings, staff))
    return settings, admin, staff, repo, employee


def test_word_excel_and_pdf_keep_real_source_locations():
    word = read_document("../safe.docx", word_bytes())
    assert word.document_name == "safe.docx"
    assert word.source_type == "docx"
    assert all(p.number is None for p in word.pages)
    assert word.pages[1].section == "교육실 안내" and word.pages[1].location == "문단 2"
    assert word.pages[2].location == "표 1" and "장소 | 교육실" in word.pages[2].text
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "물품"
    sheet.append(["물품명", "위치"])
    sheet.append(["예약 확인표", "교육실"])
    hidden = workbook.create_sheet("비공개")
    hidden.sheet_state = "hidden"
    hidden.append(["숨김 원문"])
    buffer = BytesIO()
    workbook.save(buffer)
    excel = read_document("safe.xlsx", buffer.getvalue())
    assert len(excel.pages) == 2 and excel.pages[1].location == "물품!A2:B2"
    assert excel.pages[1].number is None and "예약 확인표" in excel.pages[1].text
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, "First page")
    pdf.showPage()
    pdf.showPage()
    pdf.drawString(72, 720, "Third page")
    pdf.save()
    result = read_document("safe.pdf", buffer.getvalue())
    assert [p.number for p in result.pages if p.text] == [1, 3]


def test_excel_formula_without_cached_result_fails_closed():
    workbook = Workbook()
    workbook.active["A1"] = "=1+2"
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(PdfInputError, match="EXCEL_FORMULA"):
        read_document("formula.xlsx", buffer.getvalue())


@pytest.mark.parametrize("name,content", [("bad.docx", b"broken"), ("old.doc", b"old"), ("old.xls", b"old")])
def test_invalid_office_files_are_rejected(name, content):
    with pytest.raises(PdfInputError):
        read_document(name, content)


def test_persistent_documents_shared_after_new_session_and_restart(setup):
    settings, admin, staff, repo, employee = setup
    content = word_bytes()
    doc_id = repo.register("safe.docx", content, Model(), updated_date="2026-07-01")
    # 원본은 디스크에 보관하고 직원 세션에는 인증 정보만 둡니다.
    assert len(list((settings.library_dir / "originals").glob("*.docx"))) == 1
    fresh_auth = LocalAuth(settings.library_dir)
    fresh_auth.login("staff", "synthetic-staff-password")
    restarted = Repository(settings, fresh_auth, source_store(settings, fresh_auth))
    hits = restarted.search("교육실 예약 확인", Model().encode(["q"])[0], [doc_id], .38)
    assert hits and all(h.chunk.document_id == doc_id for h in hits)
    source = next(h.chunk for h in hits if "예약" in h.chunk.text)
    assert source.page is None and source.location == "문단 2" and source.updated_date == "2026-07-01"
    assert restarted.documents()[0]["status"] == "ready"
    assert not any(k in vars(fresh_auth) for k in ("documents", "chunks", "originals"))
    # 원본 파일을 내려받지 않고도 색인만으로 검색합니다.
    assert employee.documents()[0]["id"] == doc_id


def test_staff_cannot_bypass_hidden_controls_or_download_original(setup):
    settings, admin, staff, repo, employee = setup
    content = word_bytes()
    doc_id = repo.register("safe.docx", content, Model())
    original_key = repo._row(doc_id)["source_key"]
    entry = {"document_id": doc_id, "title": "예약", "keywords": ["예약"], "items": []}
    attempts = [
        lambda: employee.register("new.docx", word_bytes("다른 글"), Model()),
        lambda: employee.register("new.docx", word_bytes("다른 글"), Model(), replaces_id=doc_id),
        lambda: employee.retire(doc_id),
        lambda: employee.reindex(doc_id, Model()),
        lambda: employee.preview(doc_id),
        lambda: employee.documents(all_status=True),
        lambda: employee.source_chunks(doc_id),
        lambda: employee.save_checklist(entry),
        lambda: employee.store.read(original_key),
        lambda: employee.store.delete(original_key),
        lambda: staff.create_user("evil", "not-allowed-password", "admin"),
    ]
    for attempt in attempts:
        with pytest.raises(GuideError, match="관리자"):
            attempt()
    assert len(repo.documents()) == 1
    assert LocalSourceStore(settings.library_dir / "originals", admin).read(original_key) == content


def test_search_does_not_wait_for_replacement_embedding(setup):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event

    settings, admin, staff, repo, employee = setup
    doc_id = repo.register("safe.docx", word_bytes(), Model())
    building, release = Event(), Event()

    class SlowModel(Model):
        def encode(self, texts):
            building.set()
            assert release.wait(10)
            return super().encode(texts)

    with ThreadPoolExecutor(max_workers=2) as pool:
        job = pool.submit(repo.register, "updated.docx", word_bytes("새 교육실 안내입니다."),
                          SlowModel(), replaces_id=doc_id)
        assert building.wait(5)
        try:
            query = pool.submit(employee.search, "교육실", Model().encode(["q"])[0], [doc_id], .38)
            assert query.result(timeout=3)  # 색인 작업을 끝내지 않아도 기존 검색은 응답
        finally:
            release.set()
        job.result(timeout=10)


def test_replacement_is_atomic_and_failed_build_retains_old_guideline(setup):
    settings, admin, staff, repo, employee = setup
    original_id = repo.register("safe.docx", word_bytes(), Model())
    original_revision = repo.revision()

    class BrokenModel(Model):
        def encode(self, texts):
            raise RuntimeError("potential secret must not be shown")

    with pytest.raises(GuideError, match="INDEX") as error:
        repo.register("updated.docx", word_bytes("예약은 새 확인표에서 확인합니다."), BrokenModel(), replaces_id=original_id)
    assert "potential secret" not in str(error.value)
    assert repo.revision() == original_revision
    assert [d["id"] for d in employee.documents()] == [original_id]
    replacement = next(d for d in repo.documents(all_status=True) if d["status"] == "error")
    assert repo.preview(replacement["id"]).source_type == "docx"
    repo.reindex(replacement["id"], Model())
    assert [d["id"] for d in employee.documents()] == [replacement["id"]]
    assert employee.search("교육실", Model().encode(["q"])[0], [original_id], .38) == []
    repo.reindex(replacement["id"], Model())  # 교체본도 재색인 가능


def test_deletion_and_reindex_invalidate_sources_and_checklists(setup):
    settings, admin, staff, repo, employee = setup
    doc_id = repo.register("safe.docx", word_bytes(), Model())
    chunks = repo.source_chunks(doc_id)
    c = next(c for c in chunks if "예약" in c.text)
    checklist = validate_checklist("예약 확인", ["예약"], [
        {"chunk_id": c.id, "quote": c.text, "stage": "사전 확인"}], chunks)
    repo.save_checklist(checklist)
    assert employee.list_checklists([doc_id])
    revision = employee.revision()
    employee.search("교육실", Model().encode(["q"])[0], [doc_id], .38)  # warm cache
    repo.reindex(doc_id, Model())
    assert employee.list_checklists([doc_id]) == []
    assert employee.source_chunks(doc_id, [c.id]) == []
    with pytest.raises(GuideError, match="CORPUS_CHANGED"):
        employee.ensure_revision(revision)
    repo.retire(doc_id)
    assert employee.documents() == []
    assert employee.search("교육실", Model().encode(["q"])[0], [doc_id], .38) == []
    assert employee.source_chunks(doc_id, [c.id]) == []
    assert not list((settings.library_dir / "originals").glob("*.docx"))


def test_failed_original_delete_never_restores_search_visibility(setup, monkeypatch):
    settings, admin, staff, repo, employee = setup
    doc_id = repo.register("safe.docx", word_bytes(), Model())
    actual_delete = repo.store.delete
    def fail(key):
        raise GuideError("원본 삭제 재시도 (STORAGE)")
    monkeypatch.setattr(repo.store, "delete", fail)
    with pytest.raises(GuideError, match="STORAGE"):
        repo.retire(doc_id)
    assert not employee.documents()
    assert repo.documents(all_status=True)[0]["last_error"]
    monkeypatch.setattr(repo.store, "delete", actual_delete)
    repo.retire(doc_id)
    assert not repo.documents(all_status=True)[0]["last_error"]


def test_private_upload_is_rejected_before_original_storage(setup):
    settings, admin, staff, repo, employee = setup
    with pytest.raises(GuideError, match="PRIVACY"):
        repo.register("private.docx", word_bytes("환자명: 홍길동"), Model())
    assert repo.documents(all_status=True) == []
    assert not (settings.library_dir / "originals").exists()


def test_revocation_and_expired_login_block_existing_sessions(setup):
    settings, admin, staff, repo, employee = setup
    admin.deactivate(staff.user_id)
    with pytest.raises(GuideError, match="ACCESS"):
        employee.documents()
    assert not staff.token
    with pytest.raises(GuideError):
        admin.deactivate(admin.user_id)
    with admin.db() as db:
        db.execute("update sessions set expires=0")
    with pytest.raises(GuideError, match="ACCESS"):
        admin.require_admin()


def test_bootstrap_race_and_login_attempt_limit(setup):
    settings, admin, staff, repo, employee = setup
    second = LocalAuth(settings.library_dir)
    with pytest.raises(GuideError, match="이미"):
        second.bootstrap("another", "long-enough-password")
    for _ in range(5):
        with pytest.raises(GuideError, match="LOGIN"):
            second.login("staff", "incorrect")
    with pytest.raises(GuideError, match="LOGIN_LIMIT"):
        second.login("staff", "synthetic-staff-password")
    with admin.db() as db:
        row = db.execute("select salt,password_hash from users where username='admin'").fetchone()
        assert "synthetic-admin-password" not in tuple(row)


def test_storage_path_traversal_and_wrong_backend_are_rejected(setup):
    settings, admin, staff, repo, employee = setup
    with pytest.raises(GuideError, match="STORAGE_KEY"):
        repo.store.read("../accounts.sqlite3")
    doc_id = repo.register("safe.docx", word_bytes(), Model())
    wrong = Repository(replace(settings, storage_backend="supabase"), admin, repo.store)
    with pytest.raises(GuideError, match="STORAGE_CONFIG"):
        wrong.preview(doc_id)
    with pytest.raises(GuideError, match="STORAGE_CONFIG"):
        wrong.retire(doc_id)
    assert employee.documents()


def test_supabase_storage_uses_admin_jwt_and_staff_never_requests_original():
    role = "admin"
    calls = []
    key = "12345678-1234-1234-1234-123456789012.docx"
    def handle(request):
        calls.append(request)
        if request.url.path == "/auth/v1/user":
            return httpx.Response(200, json={"id": "account"})
        if request.url.path == "/rest/v1/guide_profiles":
            return httpx.Response(200, json=[{"user_id": "account", "active": True, "role": role}])
        assert request.headers["authorization"] == "Bearer user-jwt"
        assert request.headers["apikey"] == "public"
        if request.method == "GET":
            assert "/object/authenticated/guide-originals/" in request.url.path
            return httpx.Response(200, content=b"synthetic original")
        if request.method == "DELETE":
            assert json.loads(request.content) == {"prefixes": [key]}
        return httpx.Response(200, json={})
    settings = Settings(supabase_url="https://example.invalid", supabase_key="public")
    auth = StaffLibrary(settings, httpx.MockTransport(handle))
    auth.use_token({"access_token": "user-jwt", "refresh_token": "refresh", "expires_in": 3600})
    store = SupabaseSourceStore(settings, auth)
    store.put(key, b"synthetic original")
    assert store.read(key) == b"synthetic original"
    store.delete(key)
    role = "staff"
    count = sum("/storage/" in r.url.path for r in calls)
    for operation in (lambda: store.read(key), lambda: store.put(key, b"x"), lambda: store.delete(key)):
        with pytest.raises(GuideError, match="관리자"):
            operation()
    assert sum("/storage/" in r.url.path for r in calls) == count


def test_storage_sql_parses_and_defends_originals_from_staff():
    sql = (APP.parent / "storage_schema.sql").read_text(encoding="utf-8")
    assert len(parse_sql(sql)) >= 15
    assert "as restrictive for all to public" in sql
    assert "public=false" in sql and "guide_is_admin()" in sql


def test_staff_ui_has_no_admin_widgets_and_discards_revoked_conversations(setup, monkeypatch):
    settings, admin, staff, repo, employee = setup
    import streamlit as st
    st.cache_resource.clear()
    monkeypatch.setenv("GUIDE_MODE", "local")
    monkeypatch.setenv("GUIDE_DATA_DIR", str(settings.library_dir))
    monkeypatch.setenv("GUIDE_STORAGE_BACKEND", "local")
    monkeypatch.setenv("GUIDE_LLM_PROVIDER", "disabled")
    monkeypatch.setattr("mvp.library.Embedder", Model)
    doc_id = repo.register("safe.docx", word_bytes(), Model())
    app = AppTest.from_file(str(APP), default_timeout=30).run()
    assert not app.exception
    assert not app.text_area and not app.tabs
    app.session_state["auth"] = staff
    app.run()
    assert not app.exception
    assert not app.tabs  # 직원은 메뉴를 선택할 필요 없는 채팅 화면 하나만 사용
    assert not app.get("file_uploader") and not app.get("download_button")
    assert not app.multiselect
    assert not any("설정" in x.value for x in app.subheader)
    assert not any(k in app.session_state for k in ("library", "preview_document", "editor_sources"))
    assert [c.label for c in app.checkbox] == ["이전 질문에 이어서 묻기"]
    assert {b.label for b in app.button} == {"새 대화 시작", "로그아웃", "CRE 환자 격리 방법은?", "PCN irrigation 방법 알려줘", "Thoracentesis 준비물은?", "반코마이신 투여 시 주의사항은?"}
    ask(app, "교육실 예약 확인")
    assert not app.exception and app.session_state["turns"]
    # 검색은 로컬로 하고, 키 미설정 안내에도 문단 출처는 확인할 수 있습니다.
    assert any("문단 2" in e.label for e in app.expander)
    repo.retire(doc_id)
    app.run()
    assert app.session_state["turns"] == []
    admin.deactivate(staff.user_id)
    app.run()
    assert not app.tabs and not app.chat_input


def test_admin_ui_and_first_login_setup_are_separate(tmp_path, monkeypatch):
    monkeypatch.setenv("GUIDE_MODE", "local")
    monkeypatch.setenv("GUIDE_DATA_DIR", str(tmp_path / "fresh"))
    monkeypatch.setenv("GUIDE_STORAGE_BACKEND", "local")
    monkeypatch.setenv("GUIDE_LLM_PROVIDER", "disabled")
    app = AppTest.from_file(str(APP), default_timeout=30).run()
    assert not app.exception and not app.tabs
    assert any("최초 관리자" in s.value for s in app.subheader)
    app.text_input[0].set_value("admin")
    app.text_input[1].set_value("synthetic-password")
    app.text_input[2].set_value("synthetic-password")
    next(b for b in app.button if b.label == "관리자 계정 만들기").click().run()
    assert not app.exception
    assert not app.tabs
    assert {"AI 채팅", "지침서 관리"}.issubset({b.label for b in app.button})
    next(b for b in app.button if b.label == "지침서 관리").click().run()
    assert len(app.get("file_uploader")) == 1
    next(b for b in app.button if b.label == "로그아웃").click().run()
    assert not app.exception and not app.tabs
    assert any("직원 로그인" == s.value for s in app.subheader)
    assert not any("최초 관리자" in s.value for s in app.subheader)


def test_no_evidence_returns_exact_requested_message_without_ai(setup, monkeypatch):
    settings, admin, staff, repo, employee = setup
    import streamlit as st
    st.cache_resource.clear()
    monkeypatch.setenv("GUIDE_MODE", "local")
    monkeypatch.setenv("GUIDE_DATA_DIR", str(settings.library_dir))
    monkeypatch.setenv("GUIDE_STORAGE_BACKEND", "local")
    monkeypatch.setattr("mvp.library.Embedder", Model)
    monkeypatch.setattr("mvp.ai.generate", lambda *a: pytest.fail("No evidence must not call LLM"))
    repo.register("safe.docx", word_bytes(), Model())
    app = AppTest.from_file(str(APP), default_timeout=30).run()
    app.session_state["auth"] = staff
    app.run()
    ask(app, "PCN irrigation")
    assert not app.exception
    assert any(x.value == NO_GUIDELINE for x in app.markdown)
    assert NO_GUIDELINE == "등록된 지침서에서 확인할 수 없습니다."


def test_catalog_does_not_persist_staff_questions_or_answers(setup):
    settings, admin, staff, repo, employee = setup
    doc_id = repo.register("safe.docx", word_bytes(), Model())
    employee.search("고유한 검색 질문 987654", Model().encode(["q"])[0], [doc_id], .38)
    with database(repo.path) as db:
        all_sql = "\n".join(db.iterdump())
    assert "고유한 검색 질문 987654" not in all_sql


def test_admin_backup_is_consistent_verified_and_staff_is_denied(setup):
    import sqlite3

    from mvp.operations import create_backup, verify_backup
    settings, admin, staff, repo, employee = setup
    repo.register("safe.docx", word_bytes(), Model())
    with pytest.raises(GuideError):
        create_backup(employee)
    backup = create_backup(repo)
    assert verify_backup(backup)
    with sqlite3.connect(backup / 'catalog.sqlite3') as db:
        assert db.execute("select count(*) from documents where status='ready'").fetchone()[0] == 1
    with sqlite3.connect(backup / 'accounts.sqlite3') as db:
        assert db.execute('select count(*) from sessions').fetchone()[0] == 0
    assert list((backup / 'originals').glob('*.docx'))
    original = next((backup / 'originals').glob('*.docx'))
    original.write_bytes(b'tampered')
    assert not verify_backup(backup)


def test_surrounding_sources_are_revoked_when_document_changes(setup):
    settings, admin, staff, repo, employee = setup
    doc_id = repo.register("safe.docx", word_bytes(), Model())
    part = repo.source_chunks(doc_id)[0]
    revision = employee.revision()
    assert employee.source_context(part, revision)
    repo.retire(doc_id)
    with pytest.raises(GuideError, match='CORPUS_CHANGED'):
        employee.source_context(part, revision)
