"""새 검색·종합·거절 경로를 합성 문서와 네트워크 대역으로 검증합니다."""

import json
from dataclasses import replace
from io import BytesIO

import httpx
import numpy as np
import pytest
from docx import Document

from mvp.ai import Quota, generate, validate_answer
from mvp.auth import LocalAuth
from mvp.documents import read_document
from mvp.evaluate_cases import CASES, corpus, fixture_answer
from mvp.grounding import explicit_conflicts
from mvp.library import DIMENSIONS, Hit, LocalLibrary, chunk_payload, make_chunks
from mvp.query import plan_query
from mvp.repository import Repository
from mvp.settings import GuideError, Settings
from mvp.storage import source_store


class Model:
    def __init__(self):
        self.calls = []

    def count(self, text):
        return len(text)

    def encode(self, texts):
        self.calls.append(list(texts))
        vectors = np.zeros((len(texts), DIMENSIONS), dtype=np.float32)
        vectors[:, 0] = 1
        return vectors


def word_bytes(text):
    document = Document()
    document.add_heading('PCN 교육', 1)
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize('code,question,previous,expected', CASES)
def test_question_types_retrieve_expected_sources_without_dense_help(code, question, previous, expected):
    library = LocalLibrary()
    library.chunks, docs = corpus()
    library.vectors = np.zeros((len(library.chunks), DIMENSIONS), dtype=np.float32)
    plan = plan_query(question, previous=previous, documents=docs)
    hits = library.search(plan.query, np.zeros(DIMENSIONS, dtype=np.float32), [d['id'] for d in docs], .38, plan)
    assert expected.issubset({h.chunk.id for h in hits})
    if not expected:
        assert not hits
    answer = validate_answer(json.dumps(fixture_answer(hits, expected)), hits)
    assert answer.answerable == bool(expected)
    assert all(h.chunk.page and h.chunk.document_name for h in hits)
    if code == 'followup':
        assert hits[0].chunk.id == 'release'
    if code == 'typo':
        assert plan.corrections == (('irrigaton', 'irrigation'),)
    if code == 'synthesis':
        assert plan.min_documents == 2 and not plan.clarification


def test_question_planning_does_not_guess_patient_or_ambiguous_documents():
    assert plan_query('아까 질문한 환자에서는 그럼 어떻게 해야 해?', 'CRE 교육').clarification
    assert plan_query('이 두 지침 내용 차이가 뭐야?').clarification
    assert plan_query('그럼 해제 기준은?', 'CRE 격리 어떻게 해?').kind == 'release'
    assert plan_query('CPE 교육', 'CRE 교육', True).entities == ('cpe',)


def test_raw_text_location_and_adjacent_ids_survive_serialization():
    document = read_document('structure.docx', word_bytes('준비물\n교육 확인표\n\n주의사항\n실제 처방에 사용하지 않습니다.'))
    _, chunks = make_chunks(document, Model())
    for index, chunk in enumerate(chunks):
        row = chunk_payload(chunk)
        assert row['chunk_id'] == chunk.id and row['raw_text'] == chunk.text
        assert row['page_number'] is None and row['section_title'] == chunk.section
        assert row['normalized_text'] and chunk.parent_id
        assert chunk.previous_chunk_id == (chunks[index - 1].id if index else None)
        assert chunk.next_chunk_id == (chunks[index + 1].id if index + 1 < len(chunks) else None)


def test_parsing_and_embedding_reuse_on_reindex_and_replacement(tmp_path, monkeypatch):
    import mvp.documents as documents
    documents._PARSED.clear()
    original_reader, parses = documents._read_document, []
    def counted(*args, **kwargs):
        parses.append(1)
        return original_reader(*args, **kwargs)
    monkeypatch.setattr(documents, '_read_document', counted)
    data = word_bytes('PCN 교육 확인표를 읽습니다.')
    read_document('cache.docx', data)
    read_document('cache.docx', data)
    assert len(parses) == 1
    settings = Settings(mode='local', data_dir=str(tmp_path))
    auth = LocalAuth(settings.library_dir)
    auth.bootstrap('admin', 'synthetic-password')
    repo = Repository(settings, auth, source_store(settings, auth))
    model = Model()
    identifier = repo.register('cache.docx', data, model)
    first_calls = len(model.calls)
    repo.reindex(identifier, model)
    assert len(model.calls) == first_calls
    current = next(d for d in repo.documents() if d['id'] == identifier)
    assert current['embedding_computed'] == 0 and current['embedding_reused'] == current['chunk_count']
    replacement = repo.register('cache.docx', word_bytes('PCN 교육 설문지를 작성합니다.'), model, replaces_id=identifier)
    assert len(model.calls) == first_calls + 1
    assert all('PCN 교육\n' != text and text != 'PCN 교육' for text in model.calls[-1])
    assert [d['id'] for d in repo.documents()] == [replacement]


def conflict_hits():
    a = corpus()[0][0]
    return [Hit(replace(a, id='a', document_id='a', text='PCN 교육 카드 검토 시간은 10분입니다.'), .8),
            Hit(replace(a, id='b', document_id='b', text='PCN 교육 카드 검토 시간은 20분입니다.'), .8)]


def response_body(hits, conflict):
    return dict(answerable=True, conflict=conflict, statements=[dict(text=h.chunk.text,
                evidence=[dict(chunk_id=h.chunk.id, quote=h.chunk.text)]) for h in hits])


@pytest.mark.parametrize('reported_conflict', [True, False])
def test_conflicting_sources_require_both_citations_and_warning(tmp_path, reported_conflict):
    hits = conflict_hits()
    assert explicit_conflicts(hits) == [('a', 'b')]
    config = Settings(llm_provider='internal', llm_url='https://example.invalid/v1', llm_model='fixture', llm_approved=True)
    calls = []
    def response(request):
        calls.append(1)
        return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(response_body(hits, reported_conflict))}}]})
    def invoke():
        return generate(config, 'PCN 교육 카드 검토 시간', hits, 'test', Quota(tmp_path / 'quota.db'), httpx.MockTransport(response))
    if reported_conflict:
        answer, _ = invoke()
        assert answer.conflict and answer.format == 'comparison'
    else:
        with pytest.raises(GuideError, match='AI_CONFLICT'):
            invoke()
    assert len(calls) == 1


def test_comparison_missing_one_side_stops_before_network():
    hits = [Hit(corpus()[0][1], .8)]  # CRE만 제공; VRE의 근거 없음
    plan = plan_query('CRE VRE 차이를 비교해줘')
    answer, _ = generate(Settings(), plan.query, hits, 'test', plan=plan,
                         transport=httpx.MockTransport(lambda r: pytest.fail('Must not call API')))
    assert not answer.answerable


def test_streamlit_secrets_and_environment_override_dotenv(monkeypatch):
    import mvp.settings as settings
    monkeypatch.setattr(settings, 'streamlit_secrets', lambda: {'GUIDE_LLM_MODEL': 'secret-model', 'GUIDE_MODE': 'local'})
    monkeypatch.delenv('GUIDE_LLM_MODEL', raising=False)
    assert settings.load_settings().llm_model == 'secret-model'
    monkeypatch.setenv('GUIDE_LLM_MODEL', 'environment-model')
    assert settings.load_settings().llm_model == 'environment-model'


def test_room_identifier_is_blocked_before_embedding():
    with pytest.raises(GuideError, match='PRIVACY'):
        plan_query('병실번호: 501 환자 질문')


def test_presentation_step_number_is_discarded_but_clinical_label_is_checked():
    hits = [Hit(corpus()[0][0], .8)]
    payload = response_body(hits, False)
    payload['statements'][0]['label'] = 'step1'
    assert validate_answer(json.dumps(payload), hits).statements[0].label == ''
    payload['statements'][0]['label'] = '5 mg'
    with pytest.raises(GuideError, match='AI_EVIDENCE'):
        validate_answer(json.dumps(payload), hits)


def test_cited_quote_does_not_permit_invented_preparation_step():
    source = corpus()[0][1]
    payload = {'answerable': True, 'statements': [{'text': 'CRE 격리 교육 안내문을 준비한다.',
               'evidence': [{'chunk_id': source.id, 'quote': source.text}]}]}
    with pytest.raises(GuideError, match='AI_EVIDENCE'):
        validate_answer(json.dumps(payload), [Hit(source, .8)])
    payload['statements'][0]['text'] = 'CRE 교육 안내문을 읽고 학습 확인표를 확인한다.'
    # 현재 계약은 추출형 답변입니다. 의미가 유사해도 원문과 다른 문장을 표시하지 않습니다.
    with pytest.raises(GuideError, match='AI_EVIDENCE'):
        validate_answer(json.dumps(payload), [Hit(source, .8)])
    payload['statements'][0]['text'] = source.text
    assert validate_answer(json.dumps(payload), [Hit(source, .8)]).answerable
