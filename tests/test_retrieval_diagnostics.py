"""합성 문서로 차단 원인과 관리자 진단 계약을 검증합니다. 운영 문서 검증과 구분합니다."""

import csv
import json
from dataclasses import replace
from io import BytesIO

import httpx
import numpy as np
import pytest
from docx import Document

from mvp.ai import Quota, answer_text, generate
from mvp.auth import LocalAuth
from mvp.cloud import StaffLibrary
from mvp.cloud_repository import CloudRepository
from mvp.diagnostic_ui import BM25_CSV_COLUMNS, bm25_csv, bm25_debug_rows
from mvp.diagnostics import extraction_summary, run_diagnostics, safe_failure, vectors_summary
from mvp.documents import PdfDocument, PdfPage
from mvp.evidence import assess_evidence
from mvp.library import DIMENSIONS, NO_GUIDELINE, Chunk, Hit, LocalLibrary
from mvp.query import plan_query
from mvp.repository import Repository, database
from mvp.settings import GuideError, Settings
from mvp.storage import source_store

QUESTION = '진정간호 목적에 대해 알려줘'
BODY = '교육용 용어 카드를 읽어 문서 이해를 돕는다.'


def purpose_chunk(text=BODY, section='1. 목적'):
    return Chunk('purpose', 'doc', '가상 진정간호.pdf', 2, '가상 진정간호', section, None, text, 0)


class FixtureModel:
    def count(self, text):
        return len(text)

    def encode(self, texts):
        result = np.zeros((len(texts), DIMENSIONS), dtype=np.float32)
        result[:, 0] = 1
        return result


@pytest.mark.parametrize('legacy', [False, 'with_topic', 'without_topic'])
def test_purpose_body_uses_explicit_section_before_and_after_citation(tmp_path, legacy):
    prefix = '진정간호\n' if legacy == 'with_topic' else ''
    source = purpose_chunk(prefix + '1. 목적\n' + BODY, '') if legacy else purpose_chunk()
    plan, trace = plan_query(QUESTION), {}
    assert assess_evidence(plan, [Hit(source, .1)]).sufficient
    def handle(request):
        return httpx.Response(200, json={'choices': [{'message': {'content': json.dumps({
            'answerable': True, 'statements': [{'text': BODY, 'evidence': [{'chunk_id': source.id, 'quote': BODY}]}],
        })}}]})
    settings = Settings(llm_provider='internal', llm_url='https://example.invalid/v1', llm_model='fixture', llm_approved=True)
    answer, _ = generate(settings, QUESTION, [Hit(source, .1)], 'fixture', Quota(tmp_path / 'quota.db'),
                         httpx.MockTransport(handle), plan=plan, trace=trace)
    assert answer_text(answer) == BODY
    assert trace['pre_llm_assessment'] == trace['citation_assessment'] == 'supported'
    assert trace['llm_called'] and trace['block_reason'] is None


@pytest.mark.parametrize(('text', 'section', 'query'), [
    ('교육실 배포 일정을 안내합니다.', '', QUESTION),
    ('목적', '목적', QUESTION),
    (BODY, '준비물', QUESTION),
    (BODY, '목적 > 준비물', QUESTION),
    ('가상 진정간호', '목적', QUESTION),
    (BODY, '목적', '진정간호 투여량은?'),
    (BODY, '목적', '다른시술 목적에 대해 알려줘'),
])
def test_title_or_heading_alone_does_not_authorize_answer(text, section, query):
    assert not assess_evidence(plan_query(query), [Hit(purpose_chunk(text, section), .999)]).sufficient


def test_missing_topic_trace_prevents_llm():
    trace = {}
    answer, _ = generate(Settings(), QUESTION, [Hit(purpose_chunk(BODY, ''), .99)], 'fixture', trace=trace,
                         transport=httpx.MockTransport(lambda r: pytest.fail('LLM called')))
    assert answer_text(answer) == NO_GUIDELINE
    assert trace['block_reason'] == 'pre_llm:no_topic_evidence' and not trace['llm_called']


def test_trace_records_all_actual_search_stages_without_changing_results():
    source, other = purpose_chunk(), replace(purpose_chunk('관계없는 문서 내용입니다.', ''), id='other', index=1)
    library, model = LocalLibrary(), FixtureModel()
    library.chunks = [source, other]
    library.vectors = model.encode([source.text, other.text])
    vector = np.zeros(DIMENSIONS, dtype=np.float32)
    vector[0], vector[1] = .1, np.sqrt(.99)
    trace, plan = {}, plan_query(QUESTION)
    expected = library.search(QUESTION, vector, ['doc'], .38, plan)
    actual = library.search(QUESTION, vector, ['doc'], .38, plan, trace=trace)
    assert expected == actual and actual[0].chunk.id == 'purpose'
    assert trace['bm25_index_size'] == 2 and set(trace['indexed_chunk_ids']) == {'purpose', 'other'}
    assert trace['bm25_top10'][0]['score'] > 0
    assert trace['actual_query'] == QUESTION
    assert trace['bm25_top10'][0]['chunk_text'] == source.text
    assert len(trace['vector_top10']) == 2 and trace['fused_top10']
    assert trace['reranked_top5'][0]['chunk_id'] == 'purpose'
    assert trace['thresholds']['unsupported_dense_minimum'] == .55
    reasons = {d['chunk_id']: d['reason'] for d in trace['rerank_decisions']}
    assert reasons == {'purpose': 'selected', 'other': 'below_unsupported_dense_threshold'}


def test_bm25_csv_has_required_columns_top_ten_and_full_chunk_text():
    long_text = '가상 원문입니다. ' * 40
    trace = {'actual_query': '실제 확장 검색어', 'bm25_top10': [
        {'score': 8.21 - rank, 'document_name': '가상지침.pdf', 'page_number': rank + 1,
         'section_title': '가상 항목', 'chunk_id': f'chunk-{rank}', 'chunk_text': long_text + str(rank)}
        for rank in range(12)
    ]}
    rows = bm25_debug_rows('사용자 질문', trace)
    assert len(rows) == 10 and rows[0]['rank'] == 1 and rows[-1]['rank'] == 10
    assert rows[0]['query'] == '실제 확장 검색어'
    assert rows[0]['chunk_text'] == long_text + '0'
    decoded = bm25_csv('사용자 질문', trace).decode('utf-8-sig')
    parsed = list(csv.DictReader(decoded.splitlines()))
    assert tuple(parsed[0]) == BM25_CSV_COLUMNS
    assert len(parsed) == 10 and parsed[0]['bm25_score'] == '8.21'
    assert parsed[0]['chunk_text'] == long_text + '0'


def test_extraction_counts_include_empty_pages_and_whitespace_keywords():
    result = extraction_summary(PdfDocument('fixture.pdf', (PdfPage(1, ''), PdfPage(2, '진정 간호\n목적'))))
    assert result['page_count'] == 2 and result['pages'][0]['characters'] == 0
    assert result['total_characters'] == len('진정 간호\n목적')
    assert result['keywords']['진정간호'] == {'exact': 0, 'whitespace_normalized': 1}


def test_vector_audit_reports_actual_dimensions_missing_zero_and_invalid():
    valid = [1.] + [0.] * (DIMENSIONS - 1)
    rows = [('valid', json.dumps(valid)), ('missing', None), ('wrong_dim', [1, 2]),
            ('zero', [0] * DIMENSIONS), ('nan', [float('nan')] * DIMENSIONS)]
    summary, vectors = vectors_summary(rows, ['valid', 'missing_row'])
    assert summary['valid_vector_count'] == 1 and set(vectors) == {'valid'}
    assert summary['missing_embedding_ids'] == ['missing']
    assert summary['invalid_embedding_ids'] == ['wrong_dim', 'nan']
    assert summary['zero_vector_ids'] == ['zero']
    assert summary['dimensions'] == {'384': 3, '2': 1}
    assert summary['chunks_without_vector_rows'] == ['missing_row']


@pytest.fixture
def diagnostic_repo(tmp_path):
    settings = Settings(mode='local', data_dir=str(tmp_path / 'library'))
    admin = LocalAuth(settings.library_dir)
    admin.bootstrap('admin', 'synthetic-admin-password')
    admin.create_user('staff', 'synthetic-staff-password')
    repo = Repository(settings, admin, source_store(settings, admin))
    doc, buffer = Document(), BytesIO()
    doc.add_heading('1. 목적', level=1)
    doc.add_paragraph(BODY)
    doc.save(buffer)
    identifier = repo.register('가상 진정간호.docx', buffer.getvalue(), FixtureModel())
    return repo, identifier


def test_registered_document_audit_uses_saved_chunks_vectors_and_live_search(diagnostic_repo):
    repo, identifier = diagnostic_repo
    report = run_diagnostics(repo, repo.settings, FixtureModel, QUESTION, [identifier], probe_query='진정간호 목적')
    doc = report['documents'][0]
    assert doc['original_hash_matches_metadata'] is True
    assert doc['text_extraction_without_ocr']['total_characters'] > 0
    assert doc['actual_chunk_count'] == doc['vectors']['valid_vector_count'] > 0
    assert doc['active_in_document_query'] is True and doc['indexed_at']
    for trace in report['searches'].values():
        assert trace['query_embedding']['generated'] and trace['query_embedding']['dimensions'] == 384
        assert trace['bm25_index_size'] == doc['actual_chunk_count']
        assert trace['vector_top10'] and trace['stored_vector_cosine_top10']
        assert trace['evidence_assessment']['sufficient']
    assert report['final_answer'] is None and report['generation']['status'] == 'not_requested'
    assert report['errors'] == []
    assert 'synthetic-admin-password' not in json.dumps(report)


def test_retired_document_vectors_are_audited_but_not_searched(diagnostic_repo):
    repo, identifier = diagnostic_repo
    with database(repo.path) as db:
        db.execute("update documents set status='retired' where id=?", (identifier,))
        db.execute('update corpus set version=version+1')
    report = run_diagnostics(repo, repo.settings, FixtureModel, QUESTION, [identifier], call_llm=True)
    doc, trace = report['documents'][0], report['searches']['question']
    assert not doc['active_in_document_query'] and doc['actual_chunk_count'] > 0
    assert trace['allowed_document_ids'] == [] and trace['bm25_index_size'] == 0
    assert report['final_answer'] == NO_GUIDELINE and not report['generation']['llm_called']


def test_staff_cannot_use_debug_entrypoints(diagnostic_repo):
    repo, identifier = diagnostic_repo
    staff = LocalAuth(repo.settings.library_dir)
    staff.login('staff', 'synthetic-staff-password')
    repo.auth = staff
    for action in [lambda: run_diagnostics(repo, repo.settings, FixtureModel, QUESTION, [identifier]),
                   lambda: repo.diagnostic_source(identifier), lambda: repo.diagnostic_chunks(identifier),
                   lambda: repo.diagnostic_vectors(identifier),
                   lambda: repo.search(QUESTION, FixtureModel().encode(['q'])[0], [identifier], .38, trace={})]:
        with pytest.raises(GuideError, match='관리자'):
            action()


def test_safe_errors_do_not_include_raw_exception_details():
    assert safe_failure(ValueError('secret raw body')) == 'DIAGNOSTIC_STAGE_FAILED'
    assert safe_failure(GuideError('private backend detail (DATABASE/SEARCH/HTTP400/42703)')) == 'DATABASE/SEARCH/HTTP400/42703'


def cloud_diagnostic_repo():
    from dataclasses import asdict
    from hashlib import sha256

    word, buffer = Document(), BytesIO()
    word.add_heading('1. 목적', level=1)
    word.add_paragraph(BODY)
    word.save(buffer)
    content = buffer.getvalue()
    part = replace(purpose_chunk(), document_name='가상 진정간호.docx', page=None, source_type='docx')
    doc = dict(id='doc', document_name=part.document_name, title=part.title, status='active',
               file_hash=sha256(content).hexdigest(), created_at='2026-09-12T00:00:00Z',
               source_type='docx')  # indexed_at/chunk_count 부재를 성공으로 추측하지 않아야 합니다.
    state = {'role': 'admin', 'active': True, 'embedding_reads': 0, 'calls': []}
    def handle(request):
        assert request.headers.get('authorization') == 'Bearer fixture-token'
        path = request.url.path
        state['calls'].append(path)
        if path == '/auth/v1/user':
            return httpx.Response(200, json={'id': 'fixture'})
        if path == '/rest/v1/guide_profiles':
            return httpx.Response(200, json=[dict(user_id='fixture', role=state['role'], active=state['active'])])
        if path == '/rest/v1/guide_documents':
            return httpx.Response(200, json=[doc])
        if path == '/rest/v1/guide_chunks':
            columns = request.url.params['select'].split(',')
            if 'embedding' in columns:
                state['embedding_reads'] += 1
                return httpx.Response(200, json=[dict(id=part.id, embedding=json.dumps(FixtureModel().encode(['x'])[0].tolist()))])
            return httpx.Response(200, json=[{k: v for k, v in asdict(part).items() if k in columns}])
        if path == '/rest/v1/rpc/guide_search':
            return httpx.Response(200, json=[dict(asdict(part), similarity=.2)])
        pytest.fail('Unexpected mock request')
    settings = Settings(supabase_url='https://example.invalid', supabase_key='public-fixture', mode='staff')
    auth = StaffLibrary(settings, httpx.MockTransport(handle))
    auth.token, auth.expires, auth.user_id = 'fixture-token', float('inf'), 'fixture'
    class Store:
        def read(self, key):
            return content
    return CloudRepository(settings, auth, Store()), state


def test_cloud_diagnostics_reads_actual_embedding_column_and_preserves_unknown_status_fields():
    repo, state = cloud_diagnostic_repo()
    report = run_diagnostics(repo, repo.settings, FixtureModel, QUESTION, ['doc'], probe_query='진정간호 목적')
    doc, trace = report['documents'][0], report['searches']['question']
    assert report['errors'] == [] and state['embedding_reads'] == 1
    assert doc['storage_status'] == 'active' and doc['application_status'] == 'ready'
    assert doc['indexed_at'] is None and doc['stored_indexed'] is None
    assert doc['metadata_chunk_count'] is None
    assert doc['actual_chunk_count'] == doc['vectors']['valid_vector_count'] == 1
    assert trace['vector_top10'][0]['similarity'] == .2
    assert trace['stored_vector_cosine_top10'][0]['cosine'] == 1.0
    assert trace['reranked_top5'][0]['chunk_id'] == 'purpose' and trace['evidence_assessment']['sufficient']


def test_cloud_staff_cannot_access_diagnostic_vectors_or_trace():
    repo, state = cloud_diagnostic_repo()
    state['role'] = 'staff'
    for action in [lambda: run_diagnostics(repo, repo.settings, FixtureModel, QUESTION, ['doc']),
                   lambda: repo.diagnostic_vectors('doc'), lambda: repo.diagnostic_chunks('doc'),
                   lambda: repo.diagnostic_source('doc'),
                   lambda: repo.search(QUESTION, FixtureModel().encode(['q'])[0], ['doc'], .38, trace={})]:
        with pytest.raises(GuideError, match='관리자'):
            action()
    assert state['embedding_reads'] == 0
    assert '/rest/v1/guide_chunks' not in state['calls'] and '/rest/v1/rpc/guide_search' not in state['calls']


def test_admin_role_revocation_during_diagnostics_does_not_return_report():
    repo, state = cloud_diagnostic_repo()
    class RevokingModel(FixtureModel):
        def encode(self, texts):
            state['role'] = 'staff'
            return super().encode(texts)
    with pytest.raises(GuideError, match='관리자'):
        run_diagnostics(repo, repo.settings, RevokingModel, QUESTION, ['doc'])


def test_missing_original_does_not_hide_saved_chunk_retrieval():
    repo, _ = cloud_diagnostic_repo()
    class MissingStore:
        def read(self, key):
            raise GuideError('private storage response (STORAGE_READ)')
    repo.store = MissingStore()
    report = run_diagnostics(repo, repo.settings, FixtureModel, QUESTION, ['doc'])
    assert report['errors'] == [{'stage': 'doc:original_storage', 'code': 'STORAGE_READ'}]
    assert report['searches']['question']['evidence_assessment']['sufficient']
    assert 'private storage response' not in json.dumps(report)


@pytest.mark.parametrize('images', [[], [dict(name='scan')]])
def test_ocr_dependency_is_checked_only_for_scan_candidates(monkeypatch, images):
    from contextlib import nullcontext

    from mvp.pdf_layout import enhance_pdf
    text = '' if images else 'This is a text PDF containing a complete paragraph. ' * 3
    document = PdfDocument('fixture.pdf', (PdfPage(1, text),))
    class Page:
        def find_tables(self):
            return []
        def close(self):
            pass
    page = Page()
    page.images = images
    class PDF:
        pages = [page]
    calls = []
    monkeypatch.setattr('mvp.pdf_layout.bookmark_sections', lambda _: {})
    monkeypatch.setattr('pdfplumber.open', lambda _: nullcontext(PDF()))
    monkeypatch.setattr('mvp.pdf_layout.ocr_status', lambda: calls.append('checked') or (False, 'OCR unavailable'))
    monkeypatch.setattr('mvp.pdf_layout.ocr_page', lambda *a: pytest.fail('Unavailable OCR must not run'))
    result = enhance_pdf(document, b'fixture', ocr=True)
    assert result.pages[0].text == text
    if images:
        assert calls == ['checked'] and any('OCR_SETUP' in w for w in result.warnings)
    else:
        assert calls == [] and not result.warnings


def test_admin_diagnostic_form_runs_real_repository_search(monkeypatch, tmp_path):
    import streamlit as st

    from tests.test_mvp_chat import registered_app

    monkeypatch.setenv('GUIDE_MODE', 'local')
    monkeypatch.setenv('GUIDE_DATA_DIR', str(tmp_path / 'library'))
    monkeypatch.setenv('GUIDE_STORAGE_BACKEND', 'local')
    st.cache_resource.clear()
    monkeypatch.setattr('mvp.library.Embedder', FixtureModel)
    monkeypatch.setattr('mvp.ai.generate', lambda *a, **k: pytest.fail('LLM not requested'))
    app = registered_app(monkeypatch, tmp_path)
    next(b for b in app.button if b.key == 'nav_manage').click().run()
    assert not app.exception
    assert any(x.label == '진단 질문' for x in app.text_input)
    next(x for x in app.text_input if x.label == '진단 질문').set_value('교육실 사용 방법은?')
    next(b for b in app.button if b.label == '실제 검색 파이프라인 진단').click().run()
    assert not app.exception and not app.error
    assert any('BM25 top 10' in x.value for x in app.markdown)
    assert len(app.dataframe) >= 6
    # 일반 직원이 UI 경로를 강제로 지정해도 폼/원문 진단을 볼 수 없습니다.
    auth = app.session_state['auth']
    auth.create_user('diagnostic-staff', 'synthetic-staff-password')
    staff = LocalAuth(auth.path.parent)
    staff.login('diagnostic-staff', 'synthetic-staff-password')
    app.session_state['auth'] = staff
    app.session_state['ui_page'] = 'manage'
    app.run()
    assert not app.exception
    assert not any(x.label == '진단 질문' for x in app.text_input)
